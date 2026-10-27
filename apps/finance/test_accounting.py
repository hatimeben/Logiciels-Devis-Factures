from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from django.urls import reverse
import openpyxl

from apps.catalog.models import Service, ServiceCategory, TaxRate
from apps.partners.models import Customer
from .models import Invoice, InvoiceLine, Payment, PaymentAllocation
from .accounting_services import (
    get_accounting_period_range,
    get_moroccan_vat_report_data,
    get_sales_journal_data,
    get_payments_journal_data,
    build_complete_fiduciary_pack_xlsx,
)


class AccountingAndFiduciaryTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="accountant",
            password="secretpassword",
            is_staff=True,
        )
        for perm_code in ["view_invoice", "change_invoice", "view_payment", "add_payment"]:
            perm = Permission.objects.get(codename=perm_code)
            self.user.user_permissions.add(perm)
        self.client.force_login(self.user)

        self.customer = Customer.objects.create(
            name="Alpha Consulting SARL",
            ice="001234567890123",
            tax_identifier="12345678",
            city="Casablanca",
        )
        self.category = ServiceCategory.objects.create(name="Audit & Conseil")
        self.tax_20 = TaxRate.objects.create(name="TVA 20%", percentage=Decimal("20.00"), is_active=True)
        self.tax_10 = TaxRate.objects.create(name="TVA 10%", percentage=Decimal("10.00"), is_active=True)

        self.service1 = Service.objects.create(
            category=self.category,
            name="Mission Audit Financier",
            default_unit_price=Decimal("10000.00"),
        )
        self.service2 = Service.objects.create(
            category=self.category,
            name="Formation Normes IFRS",
            default_unit_price=Decimal("5000.00"),
        )

        # 1. Création de la facture en brouillon
        self.invoice = Invoice.objects.create(customer=self.customer)

        # 2. Ajout des lignes pendant que la facture est en brouillon
        # Ligne 1 : 10 000 MAD HT + 20% TVA = 2 000 TVA -> 12 000 TTC
        InvoiceLine.objects.create(
            invoice=self.invoice,
            service=self.service1,
            description="Mission Audit",
            quantity=Decimal("1.000"),
            unit_price=Decimal("10000.00"),
            discount_type="percentage",
            discount_value=Decimal("0.00"),
            tax_rate=self.tax_20,
            tax_name=self.tax_20.name,
            tax_percentage=self.tax_20.percentage,
        )
        # Ligne 2 : 5 000 MAD HT + 10% TVA = 500 TVA -> 5 500 TTC
        InvoiceLine.objects.create(
            invoice=self.invoice,
            service=self.service2,
            description="Formation",
            quantity=Decimal("1.000"),
            unit_price=Decimal("5000.00"),
            discount_type="percentage",
            discount_value=Decimal("0.00"),
            tax_rate=self.tax_10,
            tax_name=self.tax_10.name,
            tax_percentage=self.tax_10.percentage,
        )
        # Total Facture = 15 000 HT + 2 500 TVA = 17 500 TTC

        # 3. Émission de la facture au 15/10/2026
        Invoice.objects.filter(pk=self.invoice.pk).update(
            issue_date=date(2026, 10, 15),
            status=Invoice.Status.ISSUED,
            number="FAC-2026-0001",
            customer_snapshot={
                "name": self.customer.name,
                "ice": self.customer.ice,
                "tax_identifier": self.customer.tax_identifier,
            },
        )
        self.invoice.refresh_from_db()

    def test_accounting_period_range_helper(self):
        # Mois
        s, e, label = get_accounting_period_range(2026, "month", 10)
        self.assertEqual(s, date(2026, 10, 1))
        self.assertEqual(e, date(2026, 10, 31))
        self.assertIn("Octobre 2026", label)

        # Trimestre
        s, e, label = get_accounting_period_range(2026, "quarter", 4)
        self.assertEqual(s, date(2026, 10, 1))
        self.assertEqual(e, date(2026, 12, 31))
        self.assertIn("T4", label)

        # Année
        s, e, label = get_accounting_period_range(2026, "year", 1)
        self.assertEqual(s, date(2026, 1, 1))
        self.assertEqual(e, date(2026, 12, 31))

    def test_sales_journal_generation(self):
        sales = get_sales_journal_data(date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(sales["totals"]["count"], 1)
        self.assertEqual(sales["totals"]["net"], Decimal("15000.00"))
        self.assertEqual(sales["totals"]["tax"], Decimal("2500.00"))
        self.assertEqual(sales["totals"]["total"], Decimal("17500.00"))

        # Vérification des écritures en partie double
        # Débit 3421 = 17 500 | Crédit 7124 = 15 000 | Crédit 4455 = 2 500
        debits = sum(e["debit"] for e in sales["accounting_entries"])
        credits = sum(e["credit"] for e in sales["accounting_entries"])
        self.assertEqual(debits, Decimal("17500.00"))
        self.assertEqual(credits, Decimal("17500.00"))

    def test_vat_report_regime_debit(self):
        vat_debit = get_moroccan_vat_report_data(
            date(2026, 10, 1), date(2026, 10, 31), regime="debit"
        )
        self.assertEqual(vat_debit["totals"]["base_ht"], Decimal("15000.00"))
        self.assertEqual(vat_debit["totals"]["tax_amount"], Decimal("2500.00"))
        self.assertEqual(vat_debit["totals"]["total_ttc"], Decimal("17500.00"))
        self.assertEqual(len(vat_debit["summary"]), 2)  # 20% et 10%

    def test_vat_report_regime_encaissement_with_partial_payment(self):
        # Enregistrement d'un règlement partiel de 8 750 MAD (soit 50% de la facture 17 500 MAD)
        payment = Payment.objects.create(
            customer=self.customer,
            amount=Decimal("8750.00"),
            payment_date=date(2026, 10, 20),
            method=Payment.Method.BANK_TRANSFER,
            reference="VIR-BAM-9876",
        )
        PaymentAllocation.objects.create(
            payment=payment,
            invoice=self.invoice,
            amount=Decimal("8750.00"),
        )
        Invoice.objects.filter(pk=self.invoice.pk).update(status=Invoice.Status.PARTIALLY_PAID)
        self.invoice.refresh_from_db()

        # Calcul TVA au régime des encaissements sur Octobre 2026
        vat_enc = get_moroccan_vat_report_data(
            date(2026, 10, 1), date(2026, 10, 31), regime="encaissement"
        )
        # La TVA exigible doit être exactement la moitié : 1 250 MAD (50% de 2 500 MAD)
        self.assertEqual(vat_enc["totals"]["tax_amount"], Decimal("1250.00"))
        self.assertEqual(vat_enc["totals"]["base_ht"], Decimal("7500.00"))
        self.assertEqual(vat_enc["totals"]["total_ttc"], Decimal("8750.00"))

        # Vérification du relevé des paiements
        payments = get_payments_journal_data(date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(payments["totals"]["count"], 1)
        self.assertEqual(payments["totals"]["amount"], Decimal("8750.00"))

    def test_fiduciary_pack_xlsx_export(self):
        # Création d'un paiement
        payment = Payment.objects.create(
            customer=self.customer,
            amount=Decimal("17500.00"),
            payment_date=date(2026, 10, 25),
            method=Payment.Method.BANK_TRANSFER,
            reference="VIR-12345",
        )
        PaymentAllocation.objects.create(
            payment=payment,
            invoice=self.invoice,
            amount=Decimal("17500.00"),
        )

        vat_data = get_moroccan_vat_report_data(date(2026, 10, 1), date(2026, 10, 31), regime="encaissement")
        sales_data = get_sales_journal_data(date(2026, 10, 1), date(2026, 10, 31))
        payments_data = get_payments_journal_data(date(2026, 10, 1), date(2026, 10, 31))

        xlsx_bytes = build_complete_fiduciary_pack_xlsx("Octobre 2026", vat_data, sales_data, payments_data)
        self.assertTrue(len(xlsx_bytes) > 1000)

        # Chargement du classeur Excel avec openpyxl pour vérifier les onglets
        wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
        self.assertIn("Déclaration TVA & Relevé", wb.sheetnames)
        self.assertIn("Journal des Ventes", wb.sheetnames)
        self.assertIn("Journal des Règlements", wb.sheetnames)

    def test_accounting_dashboard_view_and_exports(self):
        # Page principale
        url = reverse("accounting-dashboard") + "?year=2026&period_type=month&period_value=10&regime=encaissement"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Espace Fiduciaire & Déclarations TVA")
        self.assertContains(response, "TVA Collectée Exigible")

        # Export Pack Excel
        pack_url = reverse("accounting-export-pack") + "?year=2026&period_type=month&period_value=10"
        response_pack = self.client.get(pack_url)
        self.assertEqual(response_pack.status_code, 200)
        self.assertEqual(
            response_pack["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("Pack_Fiduciaire", response_pack["Content-Disposition"])

        # Export TVA CSV
        tva_url = reverse("accounting-export-vat-csv") + "?year=2026&period_type=month&period_value=10"
        response_tva = self.client.get(tva_url)
        self.assertEqual(response_tva.status_code, 200)
        self.assertEqual(response_tva["Content-Type"], "text/csv; charset=utf-8-sig")

        # Export Ventes CSV
        sales_url = reverse("accounting-export-sales-csv") + "?year=2026&period_type=month&period_value=10"
        response_sales = self.client.get(sales_url)
        self.assertEqual(response_sales.status_code, 200)
        self.assertEqual(response_sales["Content-Type"], "text/csv; charset=utf-8-sig")

        # Export Règlements CSV
        pay_url = reverse("accounting-export-payments-csv") + "?year=2026&period_type=month&period_value=10"
        response_pay = self.client.get(pay_url)
        self.assertEqual(response_pay.status_code, 200)
        self.assertEqual(response_pay["Content-Type"], "text/csv; charset=utf-8-sig")

    def test_accounting_dashboard_ajax_refresh(self):
        """Vérifie que la requête AJAX renvoie le fragment partiel sans rechargement brutal."""
        url = reverse("accounting-dashboard") + "?year=2026&period_type=quarter&period_value=3&regime=encaissement&tab=tva"
        response = self.client.get(url, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.status_code, 200)

        # Le fragment partiel contient la synthèse et les onglets, mais pas l'en-tête HTML global
        self.assertContains(response, "Base HT Imposable")
        self.assertContains(response, "accounting-tabs")
        self.assertNotContains(response, "<!doctype html>")

    def test_accounting_dashboard_filter_quarter_and_regime_debit(self):
        """Vérifie la bonne prise en compte du trimestre et du régime des débits."""
        url = reverse("accounting-dashboard") + "?year=2026&period_type=quarter&period_value=4&regime=debit&tab=ventes"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["period_type"], "quarter")
        self.assertEqual(response.context["period_value"], 4)
        self.assertEqual(response.context["regime"], "debit")
        self.assertEqual(response.context["active_tab"], "ventes")
        self.assertIn("T4 (Octobre - Décembre) 2026", response.context["period_label"])
