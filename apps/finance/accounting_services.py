import calendar
import csv
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.db.models import Sum
from django.utils import timezone
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import CompanySettings, Invoice, Payment, PaymentAllocation
from .quote_calculations import calculate_quote_line, calculate_quote_totals


def _money(val):
    if val is None:
        val = Decimal("0.00")
    return f"{val:,.2f}".replace(",", " ").replace(".", ",")


def get_accounting_period_range(year, period_type="month", period_value=1):
    """
    Retourne (start_date, end_date, period_label) pour les déclarations comptables et TVA marocaines.
    """
    try:
        year = int(year)
    except (TypeError, ValueError):
        year = timezone.localdate().year

    month_names = [
        "", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin",
        "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre"
    ]

    if period_type == "month":
        try:
            m = max(1, min(12, int(period_value)))
        except (TypeError, ValueError):
            m = timezone.localdate().month
        last_day = calendar.monthrange(year, m)[1]
        start_date = date(year, m, 1)
        end_date = date(year, m, last_day)
        label = f"{month_names[m]} {year}"
    elif period_type == "quarter":
        try:
            q = max(1, min(4, int(period_value)))
        except (TypeError, ValueError):
            q = (timezone.localdate().month - 1) // 3 + 1
        start_month = (q - 1) * 3 + 1
        end_month = q * 3
        last_day = calendar.monthrange(year, end_month)[1]
        start_date = date(year, start_month, 1)
        end_date = date(year, end_month, last_day)
        label = f"Trimestre T{q} ({month_names[start_month]} - {month_names[end_month]}) {year}"
    else:  # year
        start_date = date(year, 1, 1)
        end_date = date(year, 12, 31)
        label = f"Année {year}"

    return start_date, end_date, label


def get_moroccan_vat_report_data(start_date, end_date, regime="encaissement"):
    """
    Calcule l'état récapitulatif de la TVA collectée selon le CGI marocain.
    - Régime des encaissements (droit commun / standard) : TVA exigible lors du règlement.
    - Régime des débits (sur option) : TVA exigible lors de la facturation.
    """
    rows = []
    summary_by_rate = {}
    order_num = 1

    total_base_ht = Decimal("0.00")
    total_tax_amount = Decimal("0.00")
    total_ttc = Decimal("0.00")

    def _add_to_summary(rate, name, base_ht, tax_amount, ttc):
        rate_key = Decimal(str(rate)) if rate is not None else Decimal("0.00")
        if rate_key not in summary_by_rate:
            summary_by_rate[rate_key] = {
                "rate": rate_key,
                "name": name or f"TVA {rate_key:.0f}%",
                "base_ht": Decimal("0.00"),
                "tax_amount": Decimal("0.00"),
                "total_ttc": Decimal("0.00"),
            }
        summary_by_rate[rate_key]["base_ht"] += base_ht
        summary_by_rate[rate_key]["tax_amount"] += tax_amount
        summary_by_rate[rate_key]["total_ttc"] += ttc

    if regime == "encaissement":
        # Encaissements perçus dans la période sur factures émises
        allocations = (
            PaymentAllocation.objects.filter(
                payment__payment_date__range=(start_date, end_date),
                invoice__status__in=[
                    Invoice.Status.ISSUED,
                    Invoice.Status.PARTIALLY_PAID,
                    Invoice.Status.PAID,
                ],
            )
            .select_related("payment", "invoice", "invoice__customer")
            .prefetch_related("invoice__lines")
            .order_by("payment__payment_date", "invoice__number", "id")
        )

        for alloc in allocations:
            inv = alloc.invoice
            payment = alloc.payment
            customer = inv.customer
            snapshot = inv.customer_snapshot or {}

            inv_ttc = inv.total_ttc
            if inv_ttc <= Decimal("0.00"):
                continue

            ratio = alloc.amount / inv_ttc

            # Ventilation par ligne de prestation selon le taux de TVA
            lines = list(inv.lines.all())
            if not lines:
                continue

            for line in lines:
                amounts = calculate_quote_line(
                    line.quantity,
                    line.unit_price,
                    line.discount_type,
                    line.discount_value,
                    line.tax_percentage,
                )
                rate = line.tax_percentage or Decimal("0.00")
                line_ht = amounts["net"] * ratio
                line_tax = amounts["tax"] * ratio
                line_ttc = amounts["total"] * ratio

                rows.append({
                    "order_num": order_num,
                    "invoice_number": inv.number or f"FAC-{inv.pk}",
                    "invoice_date": inv.issue_date or inv.created_at.date(),
                    "customer_name": snapshot.get("name") or customer.name,
                    "customer_ice": snapshot.get("ice") or getattr(customer, "ice", "") or "—",
                    "customer_if": snapshot.get("tax_identifier") or customer.tax_identifier or "—",
                    "payment_date": payment.payment_date,
                    "payment_method": payment.get_method_display(),
                    "payment_reference": payment.reference or "—",
                    "description": line.description,
                    "tax_rate": rate,
                    "tax_name": line.tax_name or ("Sans TVA" if rate == 0 else f"TVA {rate:.0f}%"),
                    "base_ht": line_ht,
                    "tax_amount": line_tax,
                    "total_ttc": line_ttc,
                })
                order_num += 1

                _add_to_summary(rate, line.tax_name, line_ht, line_tax, line_ttc)
                total_base_ht += line_ht
                total_tax_amount += line_tax
                total_ttc += line_ttc

    else:
        # Régime des débits : toutes les factures émises sur la période
        invoices = (
            Invoice.objects.filter(
                issue_date__range=(start_date, end_date),
                status__in=[
                    Invoice.Status.ISSUED,
                    Invoice.Status.PARTIALLY_PAID,
                    Invoice.Status.PAID,
                ],
            )
            .select_related("customer")
            .prefetch_related("lines")
            .order_by("issue_date", "number", "id")
        )

        for inv in invoices:
            customer = inv.customer
            snapshot = inv.customer_snapshot or {}

            for line in inv.lines.all():
                amounts = calculate_quote_line(
                    line.quantity,
                    line.unit_price,
                    line.discount_type,
                    line.discount_value,
                    line.tax_percentage,
                )
                rate = line.tax_percentage or Decimal("0.00")
                line_ht = amounts["net"]
                line_tax = amounts["tax"]
                line_ttc = amounts["total"]

                rows.append({
                    "order_num": order_num,
                    "invoice_number": inv.number or f"FAC-{inv.pk}",
                    "invoice_date": inv.issue_date,
                    "customer_name": snapshot.get("name") or customer.name,
                    "customer_ice": snapshot.get("ice") or getattr(customer, "ice", "") or "—",
                    "customer_if": snapshot.get("tax_identifier") or customer.tax_identifier or "—",
                    "payment_date": inv.issue_date,
                    "payment_method": "À l'émission (Débits)",
                    "payment_reference": "Facturation",
                    "description": line.description,
                    "tax_rate": rate,
                    "tax_name": line.tax_name or ("Sans TVA" if rate == 0 else f"TVA {rate:.0f}%"),
                    "base_ht": line_ht,
                    "tax_amount": line_tax,
                    "total_ttc": line_ttc,
                })
                order_num += 1

                _add_to_summary(rate, line.tax_name, line_ht, line_tax, line_ttc)
                total_base_ht += line_ht
                total_tax_amount += line_tax
                total_ttc += line_ttc

    sorted_summary = [summary_by_rate[k] for k in sorted(summary_by_rate.keys(), reverse=True)]

    return {
        "regime": regime,
        "start_date": start_date,
        "end_date": end_date,
        "rows": rows,
        "summary": sorted_summary,
        "totals": {
            "base_ht": total_base_ht,
            "tax_amount": total_tax_amount,
            "total_ttc": total_ttc,
            "count": len(rows),
        },
    }


def get_sales_journal_data(start_date, end_date):
    """
    Génère le journal des ventes (VE) conforme au Plan Comptable Marocain.
    """
    invoices = (
        Invoice.objects.filter(
            issue_date__range=(start_date, end_date),
            status__in=[
                Invoice.Status.ISSUED,
                Invoice.Status.PARTIALLY_PAID,
                Invoice.Status.PAID,
            ],
        )
        .select_related("customer")
        .prefetch_related("lines")
        .order_by("issue_date", "number", "id")
    )

    entries = []
    accounting_entries = []

    total_gross = Decimal("0.00")
    total_discount = Decimal("0.00")
    total_net = Decimal("0.00")
    total_tax = Decimal("0.00")
    total_ttc = Decimal("0.00")

    for inv in invoices:
        totals = calculate_quote_totals(inv.lines.all())
        snapshot = inv.customer_snapshot or {}
        cust_name = snapshot.get("name") or inv.customer.name
        cust_ice = snapshot.get("ice") or getattr(inv.customer, "ice", "") or "—"
        doc_num = inv.number or f"FAC-{inv.pk}"

        entries.append({
            "date": inv.issue_date,
            "number": doc_num,
            "customer_name": cust_name,
            "customer_ice": cust_ice,
            "gross": totals["gross"],
            "discount": totals["discount"],
            "net": totals["net"],
            "tax": totals["tax"],
            "total": totals["total"],
            "paid": inv.total_paid,
            "balance": inv.remaining_balance,
            "status_display": inv.get_status_display(),
        })

        total_gross += totals["gross"]
        total_discount += totals["discount"]
        total_net += totals["net"]
        total_tax += totals["tax"]
        total_ttc += totals["total"]

        # Écritures comptables (PCGM)
        # 1. Débit Client 3421
        accounting_entries.append({
            "date": inv.issue_date,
            "journal": "VE",
            "piece": doc_num,
            "account": "34210000",
            "account_label": f"Clients - {cust_name}",
            "libelle": f"Facture {doc_num}",
            "debit": totals["total"],
            "credit": Decimal("0.00"),
        })
        # 2. Crédit Ventes 7124
        accounting_entries.append({
            "date": inv.issue_date,
            "journal": "VE",
            "piece": doc_num,
            "account": "71240000",
            "account_label": "Prestations de services / Ventes",
            "libelle": f"Facture {doc_num}",
            "debit": Decimal("0.00"),
            "credit": totals["net"],
        })
        # 3. Crédit État TVA facturée 4455 (si TVA > 0)
        if totals["tax"] > Decimal("0.00"):
            accounting_entries.append({
                "date": inv.issue_date,
                "journal": "VE",
                "piece": doc_num,
                "account": "44550000",
                "account_label": "État, TVA facturée",
                "libelle": f"TVA s/ Facture {doc_num}",
                "debit": Decimal("0.00"),
                "credit": totals["tax"],
            })

    return {
        "start_date": start_date,
        "end_date": end_date,
        "entries": entries,
        "accounting_entries": accounting_entries,
        "totals": {
            "gross": total_gross,
            "discount": total_discount,
            "net": total_net,
            "tax": total_tax,
            "total": total_ttc,
            "count": len(entries),
        },
    }


def get_payments_journal_data(start_date, end_date):
    """
    Génère le journal des encaissements / règlements clients (TR/RG).
    """
    allocations = (
        PaymentAllocation.objects.filter(
            payment__payment_date__range=(start_date, end_date),
            invoice__status__in=[
                Invoice.Status.ISSUED,
                Invoice.Status.PARTIALLY_PAID,
                Invoice.Status.PAID,
            ],
        )
        .select_related("payment", "invoice", "invoice__customer")
        .order_by("payment__payment_date", "id")
    )

    entries = []
    accounting_entries = []
    total_amount = Decimal("0.00")

    for alloc in allocations:
        p = alloc.payment
        inv = alloc.invoice
        snapshot = inv.customer_snapshot or {}
        cust_name = snapshot.get("name") or inv.customer.name
        doc_num = inv.number or f"FAC-{inv.pk}"
        amount = alloc.amount

        entries.append({
            "date": p.payment_date,
            "payment_id": p.pk,
            "method_display": p.get_method_display(),
            "reference": p.reference or "—",
            "invoice_number": doc_num,
            "customer_name": cust_name,
            "amount": amount,
            "notes": p.notes or "",
        })

        total_amount += amount

        # Compte de trésorerie selon le mode
        account_code = "51410000" if p.method == Payment.Method.BANK_TRANSFER else "51610000"
        account_name = "Banque" if p.method == Payment.Method.BANK_TRANSFER else "Caisse / Règlements"

        # Écriture Débit Trésorerie
        accounting_entries.append({
            "date": p.payment_date,
            "journal": "TR",
            "piece": f"REG-{p.pk}",
            "account": account_code,
            "account_label": account_name,
            "libelle": f"Règlt {p.get_method_display()} {doc_num} - {cust_name}",
            "debit": amount,
            "credit": Decimal("0.00"),
        })
        # Écriture Crédit Client
        accounting_entries.append({
            "date": p.payment_date,
            "journal": "TR",
            "piece": f"REG-{p.pk}",
            "account": "34210000",
            "account_label": f"Clients - {cust_name}",
            "libelle": f"Règlt {doc_num}",
            "debit": Decimal("0.00"),
            "credit": amount,
        })

    return {
        "start_date": start_date,
        "end_date": end_date,
        "entries": entries,
        "accounting_entries": accounting_entries,
        "totals": {
            "amount": total_amount,
            "count": len(entries),
        },
    }


def build_complete_fiduciary_pack_xlsx(period_label, vat_data, sales_data, payments_data):
    """
    Génère le classeur Excel Haute Définition (.xlsx) avec 3 onglets formatés pour le fiduciaire :
    1. Déclaration TVA & Encaissements
    2. Journal des Ventes (VE)
    3. Journal des Règlements (TR)
    """
    company = CompanySettings.get_settings()
    wb = openpyxl.Workbook()

    # Styles d'excellence fiduciaire : Bleu Nuit & Or
    font_title = Font(name="Calibri", size=15, bold=True, color="0F172A")
    font_subtitle = Font(name="Calibri", size=11, bold=True, color="C59B27")
    font_company = Font(name="Calibri", size=9.5, color="475569")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=9.5, color="0F172A")
    font_bold = Font(name="Calibri", size=9.5, bold=True, color="0F172A")
    font_kpi_label = Font(name="Calibri", size=9, bold=True, color="64748B")
    font_kpi_val = Font(name="Calibri", size=13, bold=True, color="0F172A")

    fill_header = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_subtotal = PatternFill(start_color="FAF5E8", end_color="FAF5E8", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_kpi = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Side(style="thin", color="E2E8F0")
    gold_top = Side(style="medium", color="C59B27")
    cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    total_border = Border(top=gold_top, bottom=Side(style="double", color="0F172A"), left=thin_border, right=thin_border)

    money_fmt = '#,##0.00 "MAD"'
    date_fmt = "DD/MM/YYYY"

    # ==========================================
    # ONGLET 1 : DÉCLARATION DE TVA MAROCAINE
    # ==========================================
    ws_tva = wb.active
    ws_tva.title = "Déclaration TVA & Relevé"
    ws_tva.views.sheetView[0].showGridLines = True

    # Entête Entreprise
    ws_tva["A1"] = f"{company.name} — DÉCLARATION DE TVA ({'RÉGIME DES ENCAISSEMENTS' if vat_data['regime'] == 'encaissement' else 'RÉGIME DES DÉBITS'})"
    ws_tva["A1"].font = font_title
    ws_tva["A2"] = f"Période fiscale : {period_label} | {company.legal_notice()}"
    ws_tva["A2"].font = font_company

    # Blocs KPI
    kpis = [
        ("Base HT Imposable", vat_data["totals"]["base_ht"], "A4", "B5"),
        ("TVA Collectée Exigible", vat_data["totals"]["tax_amount"], "D4", "E5"),
        ("Total TTC Opérations", vat_data["totals"]["total_ttc"], "G4", "H5"),
    ]
    for label, val, top_left, bot_right in kpis:
        col1 = top_left[0]
        row1 = int(top_left[1])
        ws_tva[top_left] = label
        ws_tva[top_left].font = font_kpi_label
        ws_tva[f"{col1}{row1+1}"] = float(val)
        ws_tva[f"{col1}{row1+1}"].font = font_kpi_val
        ws_tva[f"{col1}{row1+1}"].number_format = money_fmt

    # Tableau Récapitulatif par Taux
    ws_tva["A7"] = "SYNTHÈSE DE LA TVA COLLECTÉE PAR TAUX"
    ws_tva["A7"].font = font_subtitle

    tva_headers = ["Taux de TVA", "Libellé de la Taxe", "Base Hors Taxe (MAD)", "Montant TVA (MAD)", "Total TTC (MAD)"]
    for col_idx, h in enumerate(tva_headers, start=1):
        c = ws_tva.cell(row=8, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center" if col_idx <= 2 else "right")
        c.border = cell_border

    row_idx = 9
    for item in vat_data["summary"]:
        ws_tva.cell(row=row_idx, column=1, value=f"{item['rate']:.0f} %").alignment = Alignment(horizontal="center")
        ws_tva.cell(row=row_idx, column=2, value=item["name"]).alignment = Alignment(horizontal="left")
        ws_tva.cell(row=row_idx, column=3, value=float(item["base_ht"])).number_format = money_fmt
        ws_tva.cell(row=row_idx, column=4, value=float(item["tax_amount"])).number_format = money_fmt
        ws_tva.cell(row=row_idx, column=5, value=float(item["total_ttc"])).number_format = money_fmt
        for ci in range(1, 6):
            cell = ws_tva.cell(row=row_idx, column=ci)
            cell.font = font_data
            cell.border = cell_border
        row_idx += 1

    # Total Synthèse
    ws_tva.cell(row=row_idx, column=1, value="TOTAL TVA").font = font_bold
    ws_tva.cell(row=row_idx, column=2, value="").font = font_bold
    ws_tva.cell(row=row_idx, column=3, value=float(vat_data["totals"]["base_ht"])).number_format = money_fmt
    ws_tva.cell(row=row_idx, column=4, value=float(vat_data["totals"]["tax_amount"])).number_format = money_fmt
    ws_tva.cell(row=row_idx, column=5, value=float(vat_data["totals"]["total_ttc"])).number_format = money_fmt
    for ci in range(1, 6):
        cell = ws_tva.cell(row=row_idx, column=ci)
        cell.font = font_bold
        cell.fill = fill_subtotal
        cell.border = total_border
    row_idx += 2

    # Tableau Détaillé des Opérations (Bordereau DGI)
    ws_tva.cell(row=row_idx, column=1, value="RELEVÉ DÉTAILLÉ DES OPÉRATIONS (CONFORME DGI)").font = font_subtitle
    row_idx += 1

    detail_headers = [
        "N°", "N° Facture", "Date Facture", "Client / Raison Sociale", "ICE Client", "IF Client",
        "Date Règlement", "Mode", "Référence", "Base HT (MAD)", "Taux", "TVA Collectée (MAD)", "TTC Encaissé (MAD)"
    ]
    for col_idx, h in enumerate(detail_headers, start=1):
        c = ws_tva.cell(row=row_idx, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 3, 7, 8, 11] else "right" if col_idx >= 10 else "left")
        c.border = cell_border

    row_idx += 1
    for r in vat_data["rows"]:
        ws_tva.cell(row=row_idx, column=1, value=r["order_num"]).alignment = Alignment(horizontal="center")
        ws_tva.cell(row=row_idx, column=2, value=r["invoice_number"]).alignment = Alignment(horizontal="center")
        ws_tva.cell(row=row_idx, column=3, value=r["invoice_date"].strftime("%d/%m/%Y")).alignment = Alignment(horizontal="center")
        ws_tva.cell(row=row_idx, column=4, value=r["customer_name"]).alignment = Alignment(horizontal="left")
        ws_tva.cell(row=row_idx, column=5, value=r["customer_ice"]).alignment = Alignment(horizontal="center")
        ws_tva.cell(row=row_idx, column=6, value=r["customer_if"]).alignment = Alignment(horizontal="center")
        ws_tva.cell(row=row_idx, column=7, value=r["payment_date"].strftime("%d/%m/%Y") if isinstance(r["payment_date"], date) else str(r["payment_date"])).alignment = Alignment(horizontal="center")
        ws_tva.cell(row=row_idx, column=8, value=r["payment_method"]).alignment = Alignment(horizontal="left")
        ws_tva.cell(row=row_idx, column=9, value=r["payment_reference"]).alignment = Alignment(horizontal="left")
        ws_tva.cell(row=row_idx, column=10, value=float(r["base_ht"])).number_format = money_fmt
        ws_tva.cell(row=row_idx, column=11, value=f"{r['tax_rate']:.0f} %").alignment = Alignment(horizontal="center")
        ws_tva.cell(row=row_idx, column=12, value=float(r["tax_amount"])).number_format = money_fmt
        ws_tva.cell(row=row_idx, column=13, value=float(r["total_ttc"])).number_format = money_fmt

        for ci in range(1, 14):
            c = ws_tva.cell(row=row_idx, column=ci)
            c.font = font_data
            c.border = cell_border
            if row_idx % 2 == 0:
                c.fill = fill_zebra
        row_idx += 1

    # ==========================================
    # ONGLET 2 : JOURNAL DES VENTES (VE)
    # ==========================================
    ws_ve = wb.create_sheet(title="Journal des Ventes")
    ws_ve.views.sheetView[0].showGridLines = True

    ws_ve["A1"] = f"{company.name} — JOURNAL DES VENTES (VE)"
    ws_ve["A1"].font = font_title
    ws_ve["A2"] = f"Période : {period_label} | {len(sales_data['entries'])} factures émises"
    ws_ve["A2"].font = font_company

    ve_headers = [
        "Date Pièce", "Journal", "N° Facture", "Client", "ICE Client",
        "Total Brut HT", "Remises", "Total Net HT (MAD)", "TVA (MAD)", "Total TTC (MAD)", "Réglé", "Solde Dû", "Statut"
    ]
    for col_idx, h in enumerate(ve_headers, start=1):
        c = ws_ve.cell(row=4, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 3, 5, 13] else "right" if 6 <= col_idx <= 12 else "left")
        c.border = cell_border

    row_idx = 5
    for e in sales_data["entries"]:
        ws_ve.cell(row=row_idx, column=1, value=e["date"].strftime("%d/%m/%Y")).alignment = Alignment(horizontal="center")
        ws_ve.cell(row=row_idx, column=2, value="VE").alignment = Alignment(horizontal="center")
        ws_ve.cell(row=row_idx, column=3, value=e["number"]).alignment = Alignment(horizontal="center")
        ws_ve.cell(row=row_idx, column=4, value=e["customer_name"]).alignment = Alignment(horizontal="left")
        ws_ve.cell(row=row_idx, column=5, value=e["customer_ice"]).alignment = Alignment(horizontal="center")
        ws_ve.cell(row=row_idx, column=6, value=float(e["gross"])).number_format = money_fmt
        ws_ve.cell(row=row_idx, column=7, value=float(e["discount"])).number_format = money_fmt
        ws_ve.cell(row=row_idx, column=8, value=float(e["net"])).number_format = money_fmt
        ws_ve.cell(row=row_idx, column=9, value=float(e["tax"])).number_format = money_fmt
        ws_ve.cell(row=row_idx, column=10, value=float(e["total"])).number_format = money_fmt
        ws_ve.cell(row=row_idx, column=11, value=float(e["paid"])).number_format = money_fmt
        ws_ve.cell(row=row_idx, column=12, value=float(e["balance"])).number_format = money_fmt
        ws_ve.cell(row=row_idx, column=13, value=e["status_display"]).alignment = Alignment(horizontal="center")

        for ci in range(1, 14):
            c = ws_ve.cell(row=row_idx, column=ci)
            c.font = font_data
            c.border = cell_border
            if row_idx % 2 == 0:
                c.fill = fill_zebra
        row_idx += 1

    # Totaux Ventes
    ws_ve.cell(row=row_idx, column=4, value="TOTAUX VENTES").font = font_bold
    ws_ve.cell(row=row_idx, column=6, value=float(sales_data["totals"]["gross"])).number_format = money_fmt
    ws_ve.cell(row=row_idx, column=7, value=float(sales_data["totals"]["discount"])).number_format = money_fmt
    ws_ve.cell(row=row_idx, column=8, value=float(sales_data["totals"]["net"])).number_format = money_fmt
    ws_ve.cell(row=row_idx, column=9, value=float(sales_data["totals"]["tax"])).number_format = money_fmt
    ws_ve.cell(row=row_idx, column=10, value=float(sales_data["totals"]["total"])).number_format = money_fmt
    for ci in range(1, 14):
        c = ws_ve.cell(row=row_idx, column=ci)
        c.font = font_bold
        c.fill = fill_subtotal
        c.border = total_border

    # Écritures comptables Ventes (PCGM) en dessous
    row_idx += 3
    ws_ve.cell(row=row_idx, column=1, value="ÉCRITURES COMPTABLES EN PARTIE DOUBLE (PLAN COMPTABLE MAROCAIN)").font = font_subtitle
    row_idx += 1
    ecr_headers = ["Date", "Journal", "N° Pièce", "Compte", "Intitulé du Compte", "Libellé de l'Écriture", "Débit (MAD)", "Crédit (MAD)"]
    for col_idx, h in enumerate(ecr_headers, start=1):
        c = ws_ve.cell(row=row_idx, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 3, 4] else "right" if col_idx >= 7 else "left")
        c.border = cell_border
    row_idx += 1

    total_deb = Decimal("0.00")
    total_crd = Decimal("0.00")
    for a in sales_data["accounting_entries"]:
        ws_ve.cell(row=row_idx, column=1, value=a["date"].strftime("%d/%m/%Y")).alignment = Alignment(horizontal="center")
        ws_ve.cell(row=row_idx, column=2, value=a["journal"]).alignment = Alignment(horizontal="center")
        ws_ve.cell(row=row_idx, column=3, value=a["piece"]).alignment = Alignment(horizontal="center")
        ws_ve.cell(row=row_idx, column=4, value=a["account"]).alignment = Alignment(horizontal="center")
        ws_ve.cell(row=row_idx, column=5, value=a["account_label"]).alignment = Alignment(horizontal="left")
        ws_ve.cell(row=row_idx, column=6, value=a["libelle"]).alignment = Alignment(horizontal="left")
        ws_ve.cell(row=row_idx, column=7, value=float(a["debit"])).number_format = money_fmt
        ws_ve.cell(row=row_idx, column=8, value=float(a["credit"])).number_format = money_fmt
        total_deb += a["debit"]
        total_crd += a["credit"]

        for ci in range(1, 9):
            c = ws_ve.cell(row=row_idx, column=ci)
            c.font = font_data
            c.border = cell_border
        row_idx += 1

    ws_ve.cell(row=row_idx, column=6, value="TOTAL ÉCRITURES").font = font_bold
    ws_ve.cell(row=row_idx, column=7, value=float(total_deb)).number_format = money_fmt
    ws_ve.cell(row=row_idx, column=8, value=float(total_crd)).number_format = money_fmt
    for ci in range(1, 9):
        c = ws_ve.cell(row=row_idx, column=ci)
        c.font = font_bold
        c.fill = fill_subtotal
        c.border = total_border

    # ==========================================
    # ONGLET 3 : JOURNAL DES RÈGLEMENTS (TR)
    # ==========================================
    ws_tr = wb.create_sheet(title="Journal des Règlements")
    ws_tr.views.sheetView[0].showGridLines = True

    ws_tr["A1"] = f"{company.name} — JOURNAL DES RÈGLEMENTS & ENCAISSEMENTS (TR)"
    ws_tr["A1"].font = font_title
    ws_tr["A2"] = f"Période : {period_label} | {len(payments_data['entries'])} encaissements"
    ws_tr["A2"].font = font_company

    tr_headers = ["Date Règlement", "N° Quittance", "Mode de Règlement", "Référence / N° Chèque", "Facture Concernée", "Client", "Montant Encaissé (MAD)", "Remarques"]
    for col_idx, h in enumerate(tr_headers, start=1):
        c = ws_tr.cell(row=4, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 3, 5] else "right" if col_idx == 7 else "left")
        c.border = cell_border

    row_idx = 5
    for p in payments_data["entries"]:
        ws_tr.cell(row=row_idx, column=1, value=p["date"].strftime("%d/%m/%Y")).alignment = Alignment(horizontal="center")
        ws_tr.cell(row=row_idx, column=2, value=f"REG-{p['payment_id']}").alignment = Alignment(horizontal="center")
        ws_tr.cell(row=row_idx, column=3, value=p["method_display"]).alignment = Alignment(horizontal="center")
        ws_tr.cell(row=row_idx, column=4, value=p["reference"]).alignment = Alignment(horizontal="left")
        ws_tr.cell(row=row_idx, column=5, value=p["invoice_number"]).alignment = Alignment(horizontal="center")
        ws_tr.cell(row=row_idx, column=6, value=p["customer_name"]).alignment = Alignment(horizontal="left")
        ws_tr.cell(row=row_idx, column=7, value=float(p["amount"])).number_format = money_fmt
        ws_tr.cell(row=row_idx, column=8, value=p["notes"]).alignment = Alignment(horizontal="left")

        for ci in range(1, 9):
            c = ws_tr.cell(row=row_idx, column=ci)
            c.font = font_data
            c.border = cell_border
            if row_idx % 2 == 0:
                c.fill = fill_zebra
        row_idx += 1

    ws_tr.cell(row=row_idx, column=6, value="TOTAL ENCAISSEMENTS").font = font_bold
    ws_tr.cell(row=row_idx, column=7, value=float(payments_data["totals"]["amount"])).number_format = money_fmt
    for ci in range(1, 9):
        c = ws_tr.cell(row=row_idx, column=ci)
        c.font = font_bold
        c.fill = fill_subtotal
        c.border = total_border

    # Auto-ajustement de la largeur des colonnes sur tous les onglets
    for sheet in [ws_tva, ws_ve, ws_tr]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2]:
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv_export(headers, rows_dict_list, key_order):
    """
    Export CSV encodé en UTF-8 avec BOM (utf-8-sig) et séparateur point-virgule (;)
    pour ouverture directe et native dans Microsoft Excel sans assistant d'import.
    """
    buffer = BytesIO()
    # Écriture du BOM UTF-8
    buffer.write(b"\xef\xbb\xbf")
    writer_wrapper = BytesIO()

    # On utilise un buffer texte temporaire
    lines = []
    lines.append(";".join(f'"{h}"' for h in headers))
    for r in rows_dict_list:
        row_vals = []
        for k in key_order:
            val = r.get(k, "")
            if isinstance(val, (int, float, Decimal)):
                val_str = f"{val:.2f}".replace(".", ",")
            elif isinstance(val, date):
                val_str = val.strftime("%d/%m/%Y")
            else:
                val_str = str(val).replace('"', '""')
            row_vals.append(f'"{val_str}"')
        lines.append(";".join(row_vals))

    return "\r\n".join(lines).encode("utf-8-sig")

